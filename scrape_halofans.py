#!/usr/bin/env python3
"""
Halofans / Moflip event scraper
================================

Scrapes events from https://vesta.halofans.id, which is a Next.js front-end
backed by TWO different JSON APIs:

  1. api.halofans.id  -> older, numeric-id events (/event/{id})
  2. spl.moflip.com   -> newer "v2" slug events   (/event/v2/{slug})

Nothing here needs a browser: every command talks to the JSON APIs directly.

--------------------------------------------------------------------------
DISCOVERED ENDPOINTS
--------------------------------------------------------------------------
  Legacy (api.halofans.id):
    GET /v1/search?q={kw}&filter[]=Event&...&page={n}&per_page={n}
    GET /v1/events/{numeric_id}

  v2 / moflip (spl.moflip.com):
    GET /events?limit={n}&offset={n}       -> paginated listing (event ids)
    GET /events/{slug_or_id}               -> full detail + tickets
    GET /events/{slug}/ticket-status       -> live availability

--------------------------------------------------------------------------
COMMANDS  (run `python scrape_halofans.py -h` for the full list)
--------------------------------------------------------------------------
  list-v2            List every v2 event and resolve all slugs.
      --details      ...also fetch full detail + tickets for each.
  v2 SLUG [SLUG...]  Scrape specific v2 slug events (e.g. indo-comic).
  extract-codes      Find invitation / compliment codes across all v2 events.
  search KEYWORD     Search legacy events by keyword.
      --pages N
  ids ID [ID...]     Fetch specific legacy numeric event ids.
  crawl              Sweep keywords to harvest legacy event ids.
      --details

All output is written to ./output/ (JSON + CSV).

--------------------------------------------------------------------------
EXAMPLES
--------------------------------------------------------------------------
  python scrape_halofans.py list-v2 --details      # everything, the usual run
  python scrape_halofans.py v2 indo-comic
  python scrape_halofans.py extract-codes          # find compliment codes
  python scrape_halofans.py search "psim" --pages 2
"""

import argparse
import csv
import json
import sys
import time
from pathlib import Path
from urllib.parse import urlencode

try:
    import requests
except ImportError:
    sys.exit("Missing dependency. Run:  pip install requests")

# --------------------------------------------------------------------------- #
# Config
# --------------------------------------------------------------------------- #
import os

SITE_BASE = "https://vesta.halofans.id"
API_BASE = "https://api.halofans.id"      # legacy numeric-id events
SPL_BASE = "https://spl.moflip.com"       # v2 slug events

OUT_DIR = Path(__file__).parent / "output"

# --- Telegram notifications -------------------------------------------------
# Set these via environment variables (recommended) OR edit the fallback below.
#   export TELEGRAM_BOT_TOKEN="123456:ABC-DEF..."
#   export TELEGRAM_CHAT_ID="123456789"
# You can also pass --telegram-token / --telegram-chat on the command line.
TELEGRAM_BOT_TOKEN = os.environ.get("TELEGRAM_BOT_TOKEN", "")
TELEGRAM_CHAT_ID = os.environ.get("TELEGRAM_CHAT_ID", "")

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Accept": "application/json",
    "Referer": SITE_BASE + "/",
    "Origin": SITE_BASE,
}

SEARCH_FILTERS = ["Event", "Collectible", "Games", "Merch", "Artist"]

# Words in a name that mark an event as internal / test (not a real public event).
TEST_MARKERS = ("test", "load test", "do not buy", "internal", "[testing]",
                "buat test", "-copy-", "dummy", "sandbox")


def is_test_event(name):
    """True if an event name looks like an internal / test event."""
    n = (name or "").lower()
    return any(m in n for m in TEST_MARKERS)


def session():
    s = requests.Session()
    s.headers.update(HEADERS)
    return s


def _get_json(sess, url):
    """GET a URL and return parsed JSON, or None on any failure."""
    try:
        r = sess.get(url, timeout=30)
    except requests.RequestException as e:
        print(f"  ! request failed: {url} ({e})", file=sys.stderr)
        return None
    if r.status_code != 200:
        return None
    try:
        return r.json()
    except ValueError:
        return None


def _write(path, obj):
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")


# =========================================================================== #
# v2 / moflip events (the main, current system)
# =========================================================================== #
def list_all_v2_events(sess, limit=20, delay=0.2):
    """Return raw listing objects for every v2 event (walks all pages)."""
    events, offset = [], 0
    while True:
        j = _get_json(sess, f"{SPL_BASE}/events?limit={limit}&offset={offset}")
        if not j:
            break
        data = j.get("data", [])
        events.extend(data)
        pag = j.get("meta", {}).get("pagination", {})
        total = pag.get("total", len(events))
        if len(events) >= total or not data:
            break
        offset += limit
        time.sleep(delay)
    # dedupe by id, preserve order
    seen, out = set(), []
    for e in events:
        if e.get("id") not in seen:
            seen.add(e.get("id"))
            out.append(e)
    return out


def fetch_v2_event(sess, slug_or_id, delay=0.3):
    """Fetch a v2 event's detail + ticket-status. Returns a merged bundle."""
    detail = _get_json(sess, f"{SPL_BASE}/events/{slug_or_id}")
    time.sleep(delay)
    if not detail:
        print(f"  ! v2 event '{slug_or_id}' not found", file=sys.stderr)
        return None
    ev = detail.get("data", {}).get("moflip_event", {})
    slug = ev.get("slug", slug_or_id)
    status = _get_json(sess, f"{SPL_BASE}/events/{slug}/ticket-status")
    time.sleep(delay)
    return {
        "slug": slug,
        "page_url": f"{SITE_BASE}/event/v2/{slug}",
        "detail": detail,
        "ticket_status": status,
    }


def summarize_v2(bundle):
    """Compact, human-readable summary of a v2 event bundle."""
    ev = bundle["detail"]["data"]["moflip_event"]
    meta = ev.get("meta", {})
    tickets = []
    for t in ev.get("moflip_tickets", []):
        tickets.append({
            "id": t.get("id"),
            "name": t.get("name"),
            "price": t.get("price"),
            "sale_start": t.get("sale_start"),
            "sale_end": t.get("sale_end"),
            "availability": t.get("status", "N/A"),
        })
    tickets.sort(key=lambda x: (x["price"], x["id"]))
    return {
        "id": ev.get("id"),
        "slug": ev.get("slug"),
        "name": ev.get("name"),
        "status": ev.get("status"),
        "date": meta.get("date"),
        "location": meta.get("location"),
        "description": meta.get("desc"),
        "banner": meta.get("banner"),
        "event_start": ev.get("event_start"),
        "event_end": ev.get("event_end"),
        "page_url": bundle["page_url"],
        "invitation_code": extract_invitation(ev),
        "tickets": tickets,
    }


def extract_invitation(ev):
    """Pull invitation / compliment code config from an event, if any."""
    meta = ev.get("meta", {})
    ci = meta.get("custom_invitation")
    if not isinstance(ci, dict) or not ci:
        return None
    return {
        "field": meta.get("invitation_code_identifier"),
        "match_type": ci.get("type"),
        "valid_values": ci.get("value"),
        "min_length": ci.get("min_length"),
        "max_length": ci.get("max_length"),
    }


def _save_v2_details(sess, slugs):
    """Fetch full detail for each slug, write raw dumps + summary + ticket CSV."""
    summaries = []
    for slug in slugs:
        bundle = fetch_v2_event(sess, slug)
        if not bundle:
            continue
        _write(OUT_DIR / f"v2_{slug}.json", bundle)
        summaries.append(summarize_v2(bundle))
    if not summaries:
        return summaries
    _write(OUT_DIR / "v2_summary.json", summaries)
    with (OUT_DIR / "v2_tickets.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["event_id", "event_name", "slug", "page_url", "ticket_id",
                    "ticket_name", "price", "availability", "sale_start", "sale_end"])
        for s in summaries:
            for t in s["tickets"]:
                w.writerow([s["id"], s["name"], s["slug"], s["page_url"], t["id"],
                            t["name"], t["price"], t["availability"],
                            t["sale_start"], t["sale_end"]])
    return summaries


def cmd_list_v2(args):
    sess = session()
    print("Listing all v2 events from spl.moflip.com ...\n")
    listing = list_all_v2_events(sess)

    rows = []
    for e in listing:
        detail = _get_json(sess, f"{SPL_BASE}/events/{e['id']}")
        time.sleep(0.15)
        ev = (detail or {}).get("data", {}).get("moflip_event", {})
        slug = ev.get("slug")
        rows.append({
            "id": e["id"],
            "slug": slug,
            "name": ev.get("name") or e.get("name"),
            "status": ev.get("status") or e.get("status"),
            "event_start": e.get("event_start"),
            "event_end": e.get("event_end"),
            "url": f"{SITE_BASE}/event/v2/{slug}" if slug else None,
            "api_url": f"{SPL_BASE}/events/{slug or e['id']}",
        })
        print(f"  {e['id']:6}  {str(slug):38}  {rows[-1]['name']}")

    _write(OUT_DIR / "v2_all_slugs.json", rows)
    with (OUT_DIR / "v2_all_slugs.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["id", "slug", "name", "status",
                                          "event_start", "event_end", "url", "api_url"])
        w.writeheader()
        w.writerows(rows)

    with_slug = sum(1 for r in rows if r["slug"])
    print(f"\nFound {len(rows)} v2 events ({with_slug} with slugs).")
    print(f"  {OUT_DIR/'v2_all_slugs.json'}")
    print(f"  {OUT_DIR/'v2_all_slugs.csv'}")

    if args.details:
        slugs = [r["slug"] for r in rows if r["slug"]]
        print(f"\nFetching full detail for {len(slugs)} events...")
        summaries = _save_v2_details(sess, slugs)
        print(f"Saved details for {len(summaries)} events -> v2_summary.json, v2_tickets.csv")


def cmd_v2(args):
    sess = session()
    summaries = _save_v2_details(sess, args.slugs)
    for s in summaries:
        print(f"\n=== {s['name']} (id={s['id']}, slug={s['slug']}) ===")
        print(f"  {s['date']}  @ {s['location']}")
        print(f"  {s['page_url']}")
        if s["invitation_code"]:
            print(f"  INVITATION CODE -> {s['invitation_code']}")
        print(f"  tickets ({len(s['tickets'])}):")
        for t in s["tickets"]:
            price = "FREE" if not t["price"] else f"Rp{t['price']:,}"
            print(f"    - {t['name']:42s} {price:>12}  [{t['availability']}]")
    if summaries:
        print(f"\nSaved -> {OUT_DIR}/v2_<slug>.json, v2_summary.json, v2_tickets.csv")


def cmd_extract_codes(args):
    """Scan every v2 event for invitation / compliment codes."""
    sess = session()
    print("Scanning all v2 events for invitation / compliment codes ...\n")
    listing = list_all_v2_events(sess)
    findings = []
    for e in listing:
        detail = _get_json(sess, f"{SPL_BASE}/events/{e['id']}")
        time.sleep(0.15)
        ev = (detail or {}).get("data", {}).get("moflip_event", {})
        inv = extract_invitation(ev)
        if not inv:
            continue
        tickets = ev.get("moflip_tickets", [])
        all_free = bool(tickets) and all((t.get("price") == 0) for t in tickets)
        findings.append({
            "id": ev.get("id"),
            "slug": ev.get("slug"),
            "name": ev.get("name"),
            "all_tickets_free": all_free,
            "page_url": f"{SITE_BASE}/event/v2/{ev.get('slug')}",
            **inv,
        })

    _write(OUT_DIR / "v2_invitation_codes.json", findings)
    with (OUT_DIR / "v2_invitation_codes.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["id", "slug", "name", "all_tickets_free",
                                          "field", "match_type", "valid_values",
                                          "min_length", "max_length", "page_url"])
        w.writeheader()
        for r in findings:
            row = dict(r)
            row["valid_values"] = ", ".join(map(str, r.get("valid_values") or []))
            w.writerow(row)

    print(f"{'SLUG':38s} {'FREE?':6s} FIELD    CODES")
    print("-" * 90)
    for r in findings:
        free = "FREE" if r["all_tickets_free"] else "paid"
        print(f"{str(r['slug']):38s} {free:6s} {str(r['field']):8s} {r['valid_values']}")

    comp = [r for r in findings if r["all_tickets_free"]]
    print(f"\n{len(findings)} events have code config; "
          f"{len(comp)} are FREE compliment events:")
    for r in comp:
        print(f"  • {r['name']}")
        print(f"    slug : {r['slug']}   code(s): {r['valid_values']}")
        print(f"    page : {r['page_url']}")
    print(f"\nSaved -> {OUT_DIR/'v2_invitation_codes.json'} / .csv")


# =========================================================================== #
# Legacy numeric-id events (api.halofans.id)
# =========================================================================== #
def _search(sess, keyword, page, per_page):
    params = [("q", keyword), ("categories", "")]
    params += [("filter[]", fil) for fil in SEARCH_FILTERS]
    params += [("page", page), ("per_page", per_page)]
    return _get_json(sess, f"{API_BASE}/v1/search?{urlencode(params)}")


def fetch_legacy_event(sess, event_id, delay=0.3):
    j = _get_json(sess, f"{API_BASE}/v1/events/{event_id}")
    time.sleep(delay)
    if not j or j.get("error"):
        print(f"  ! legacy event {event_id} not found", file=sys.stderr)
        return None
    return j.get("data")


LEGACY_CSV = ["id", "title", "date", "start_date", "end_date", "location",
              "sale_start", "transaction_source", "status", "url", "short_desc"]


def _flatten_legacy(ev):
    return {
        "id": ev.get("id"), "title": ev.get("title"), "date": ev.get("date"),
        "start_date": ev.get("start_date"), "end_date": ev.get("end_date"),
        "location": ev.get("location"), "sale_start": ev.get("sale_start"),
        "transaction_source": ev.get("transaction_source"), "status": ev.get("status"),
        "url": f"{SITE_BASE}/event/{ev.get('id')}",
        "short_desc": (ev.get("short_desc") or "").replace("\n", " ").strip()[:300],
    }


def _save_legacy(events):
    _write(OUT_DIR / "events.json", events)
    with (OUT_DIR / "events.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=LEGACY_CSV)
        w.writeheader()
        for ev in events:
            w.writerow(_flatten_legacy(ev))
    print(f"\nSaved {len(events)} events -> {OUT_DIR/'events.json'}, {OUT_DIR/'events.csv'}")


def cmd_search(args):
    sess = session()
    seen, ids = set(), []
    for page in range(1, args.pages + 1):
        j = _search(sess, args.keyword, page, args.per_page)
        results = (j or {}).get("data", {}).get("event", {}).get("result", []) if j else []
        if not results:
            break
        for e in results:
            if e["id"] not in seen:
                seen.add(e["id"])
                ids.append(e["id"])
        print(f"  '{args.keyword}' page {page}: {len(results)} results (total {len(ids)})")
        time.sleep(0.3)
    print(f"\nFetching details for {len(ids)} events...")
    details = [d for d in (fetch_legacy_event(sess, i) for i in ids) if d]
    _save_legacy(details)


def cmd_ids(args):
    sess = session()
    details = [d for d in (fetch_legacy_event(sess, i) for i in args.ids) if d]
    _save_legacy(details)


CRAWL_SEEDS = (
    list("abcdefghijklmnopqrstuvwxyz0123456789")
    + ["indonesia", "jakarta", "yogyakarta", "bandung", "surabaya", "konser",
       "festival", "fan", "liga", "united", "vs", "cup", "concert", "music",
       "art", "comic", "con", "the", "and", "of", "on", "for", "night"]
)


def cmd_crawl(args):
    sess = session()
    print("Sweeping search API for all discoverable legacy event ids...\n")
    found = {}
    for kw in CRAWL_SEEDS:
        page = 1
        while page <= args.max_pages:
            j = _search(sess, kw, page, args.per_page)
            ev = (j or {}).get("data", {}).get("event", {}) if j else {}
            results = ev.get("result", [])
            if not results:
                break
            for e in results:
                found[e["id"]] = e.get("title")
            total_pages = ev.get("meta", {}).get("pagination", {}).get("total_pages") or 1
            time.sleep(0.25)
            if page >= total_pages:
                break
            page += 1
        print(f"  '{kw}' -> total {len(found)} unique events")

    rows = [{"id": i, "title": t, "url": f"{SITE_BASE}/event/{i}"}
            for i, t in sorted(found.items())]
    _write(OUT_DIR / "all_event_ids.json", rows)
    with (OUT_DIR / "all_event_ids.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["id", "title", "url"])
        w.writeheader()
        w.writerows(rows)
    print(f"\nDiscovered {len(rows)} unique legacy events -> all_event_ids.json / .csv")

    if args.details:
        details = [d for d in (fetch_legacy_event(sess, r["id"]) for r in rows) if d]
        _save_legacy(details)


# =========================================================================== #
# WATCH — detect new events / new compliment codes since last run
# =========================================================================== #
def collect_v2_snapshot(sess):
    """Return a dict snapshot of all v2 events keyed by id, incl. code info."""
    listing = list_all_v2_events(sess)
    snap = {}
    for e in listing:
        detail = _get_json(sess, f"{SPL_BASE}/events/{e['id']}")
        time.sleep(0.15)
        ev = (detail or {}).get("data", {}).get("moflip_event", {})
        slug = ev.get("slug")
        inv = extract_invitation(ev)
        tickets = ev.get("moflip_tickets", [])
        all_free = bool(tickets) and all((t.get("price") == 0) for t in tickets)
        snap[str(e["id"])] = {
            "id": e["id"],
            "slug": slug,
            "name": ev.get("name") or e.get("name"),
            "status": ev.get("status") or e.get("status"),
            "is_test": is_test_event(ev.get("name") or e.get("name")),
            "all_tickets_free": all_free,
            "codes": (inv or {}).get("valid_values"),
            "code_field": (inv or {}).get("field"),
            "url": f"{SITE_BASE}/event/v2/{slug}" if slug else None,
        }
    return snap


def _notify(title, message):
    """Best-effort desktop notification (macOS/Linux). Silently ignored if unavailable."""
    import shutil
    import subprocess
    try:
        if sys.platform == "darwin":
            script = f'display notification "{message}" with title "{title}"'
            subprocess.run(["osascript", "-e", script], check=False)
        elif shutil.which("notify-send"):
            subprocess.run(["notify-send", title, message], check=False)
    except Exception:
        pass


def _html_escape(s):
    return (str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def build_telegram_message(new_events, new_codes):
    """Build an HTML Telegram message with clickable links + visible codes."""
    parts = ["<b>🔔 Halofans — ada yang baru!</b>"]
    if new_events:
        parts.append(f"\n<b>🆕 {len(new_events)} event baru:</b>")
        for r in new_events:
            free = " 🆓 <b>FREE</b>" if r["all_tickets_free"] else ""
            name = _html_escape(r["name"])
            url = r["url"]
            line = f'• <a href="{url}">{name}</a>{free}'
            if r.get("codes"):
                codes = ", ".join(f"<code>{_html_escape(c)}</code>" for c in r["codes"])
                line += f"\n   🎟️ code: {codes}"
            parts.append(line)
    if new_codes:
        parts.append(f"\n<b>🎟️ {len(new_codes)} kode baru/berubah:</b>")
        for r in new_codes:
            name = _html_escape(r["name"])
            codes = ", ".join(f"<code>{_html_escape(c)}</code>" for c in (r.get("codes") or []))
            parts.append(f'• <a href="{r["url"]}">{name}</a>\n   🎟️ {codes}')
    return "\n".join(parts)


def send_telegram(token, chat_id, html_message):
    """Send an HTML message to Telegram. Returns True on success."""
    if not token or not chat_id:
        print("  ! Telegram tidak dikirim: token/chat_id belum di-set.", file=sys.stderr)
        return False
    url = f"https://api.telegram.org/bot{token}/sendMessage"
    try:
        r = requests.post(url, timeout=30, data={
            "chat_id": chat_id,
            "text": html_message,
            "parse_mode": "HTML",
            "disable_web_page_preview": "false",
        })
    except requests.RequestException as e:
        print(f"  ! Telegram error: {e}", file=sys.stderr)
        return False
    if r.status_code != 200:
        print(f"  ! Telegram HTTP {r.status_code}: {r.text[:200]}", file=sys.stderr)
        return False
    return True


def cmd_watch(args):
    sess = session()
    state_path = OUT_DIR / "watch_state.json"
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    token = args.telegram_token or TELEGRAM_BOT_TOKEN
    chat_id = args.telegram_chat or TELEGRAM_CHAT_ID

    # Quick connectivity test, then exit.
    if args.test_telegram:
        ok = send_telegram(token, chat_id,
                           "<b>✅ Test dari scraper Halofans</b>\nNotifikasi Telegram aktif!")
        print("Telegram test terkirim ✅" if ok else "Telegram test GAGAL ❌ (cek token/chat id)")
        return

    old = {}
    if state_path.exists():
        try:
            old = json.loads(state_path.read_text(encoding="utf-8"))
        except ValueError:
            old = {}

    print("Checking spl.moflip.com for changes ...")
    new = collect_v2_snapshot(sess)

    include = (lambda r: True) if args.include_test else (lambda r: not r["is_test"])

    new_events, new_codes = [], []
    for eid, cur in new.items():
        if not include(cur):
            continue
        prev = old.get(eid)
        if prev is None:
            new_events.append(cur)
            if cur["codes"]:
                new_codes.append(cur)
        else:
            # event existed before — did a code appear/change?
            if cur["codes"] and cur["codes"] != prev.get("codes"):
                new_codes.append(cur)

    # Always refresh the stored state (store the full new snapshot).
    state_path.write_text(json.dumps(new, ensure_ascii=False, indent=2), encoding="utf-8")

    first_run = not old
    if first_run:
        print(f"\nFirst run — baseline saved ({len(new)} events). "
              f"Run again later to see what's new.")
        return

    if not new_events and not new_codes:
        print("\nNo new events or codes since last run. ✅")
        return

    lines = []
    if new_events:
        lines.append(f"\n🆕 {len(new_events)} NEW EVENT(S):")
        for r in new_events:
            free = " [FREE]" if r["all_tickets_free"] else ""
            lines.append(f"  • {r['name']}{free}")
            lines.append(f"    {r['url']}")
            if r["codes"]:
                lines.append(f"    code(s): {', '.join(map(str, r['codes']))}")
    if new_codes:
        lines.append(f"\n🎟️  {len(new_codes)} NEW/CHANGED CODE(S):")
        for r in new_codes:
            lines.append(f"  • {r['name']}  ->  {', '.join(map(str, r['codes']))}")
            lines.append(f"    {r['url']}")
    report = "\n".join(lines)
    print(report)

    # write a report file
    (OUT_DIR / "watch_report.txt").write_text(report.strip(), encoding="utf-8")

    # optional desktop notification
    if args.notify:
        summary = f"{len(new_events)} event baru, {len(new_codes)} kode baru"
        _notify("Halofans: ada yang baru!", summary)

    # Telegram notification (default ON if token+chat available; disable with --no-telegram)
    if not args.no_telegram and (token and chat_id):
        msg = build_telegram_message(new_events, new_codes)
        if send_telegram(token, chat_id, msg):
            print("\nNotifikasi Telegram terkirim ✅")

    print(f"\nReport saved -> {OUT_DIR/'watch_report.txt'}")


# =========================================================================== #
# EXPORT — write everything to one Excel workbook (falls back to CSV)
# =========================================================================== #
def cmd_export_excel(args):
    """Combine v2 events, tickets and codes into one .xlsx (needs openpyxl)."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        sys.exit("Butuh openpyxl. Jalankan:  pip install openpyxl\n"
                 "(File CSV di folder output/ tetap bisa dibuka langsung di Excel.)")

    sess = session()
    print("Collecting all v2 events for Excel export ...")
    listing = list_all_v2_events(sess)
    slugs, base_rows = [], []
    for e in listing:
        detail = _get_json(sess, f"{SPL_BASE}/events/{e['id']}")
        time.sleep(0.15)
        ev = (detail or {}).get("data", {}).get("moflip_event", {})
        slug = ev.get("slug")
        if slug:
            slugs.append(slug)
        inv = extract_invitation(ev)
        tickets = ev.get("moflip_tickets", [])
        base_rows.append({
            "id": e["id"], "slug": slug, "name": ev.get("name") or e.get("name"),
            "status": ev.get("status") or e.get("status"),
            "is_test": is_test_event(ev.get("name") or e.get("name")),
            "event_start": e.get("event_start"), "event_end": e.get("event_end"),
            "all_free": bool(tickets) and all(t.get("price") == 0 for t in tickets),
            "codes": ", ".join(map(str, (inv or {}).get("valid_values") or [])),
            "url": f"{SITE_BASE}/event/v2/{slug}" if slug else None,
        })

    print(f"Fetching ticket detail for {len(slugs)} events ...")
    summaries = _save_v2_details(sess, slugs)
    sum_by_slug = {s["slug"]: s for s in summaries}

    wb = Workbook()
    bold = Font(bold=True)

    def sheet(ws, headers, rows):
        ws.append(headers)
        for c in ws[1]:
            c.font = bold
        for r in rows:
            ws.append(r)

    ws1 = wb.active
    ws1.title = "Events"
    sheet(ws1, ["ID", "Slug", "Name", "Status", "Test?", "Start", "End",
                "All Free", "Codes", "URL"],
          [[r["id"], r["slug"], r["name"], r["status"], "YES" if r["is_test"] else "",
            r["event_start"], r["event_end"], "YES" if r["all_free"] else "",
            r["codes"], r["url"]] for r in base_rows])

    ws2 = wb.create_sheet("Tickets")
    trows = []
    for r in base_rows:
        s = sum_by_slug.get(r["slug"])
        if not s:
            continue
        for t in s["tickets"]:
            trows.append([s["id"], s["name"], s["slug"],
                          t["name"], t["price"] or 0, t["availability"],
                          t["sale_start"], t["sale_end"]])
    sheet(ws2, ["Event ID", "Event", "Slug", "Ticket", "Price (Rp)",
                "Availability", "Sale Start", "Sale End"], trows)

    ws3 = wb.create_sheet("Codes")
    crows = [[r["id"], r["name"], r["slug"], "YES" if r["all_free"] else "",
              r["codes"], r["url"]]
             for r in base_rows if r["codes"]]
    sheet(ws3, ["Event ID", "Name", "Slug", "All Free", "Code(s)", "URL"], crows)

    # auto-ish column widths
    for ws in (ws1, ws2, ws3):
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(width + 2, 60)

    out = OUT_DIR / "halofans_events.xlsx"
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"\nSaved Excel workbook -> {out}")
    print(f"  Events : {len(base_rows)} rows")
    print(f"  Tickets: {len(trows)} rows")
    print(f"  Codes  : {len(crows)} rows")


# =========================================================================== #
# CLI
# =========================================================================== #
def main():
    p = argparse.ArgumentParser(
        description="Scrape events from vesta.halofans.id (api.halofans.id + spl.moflip.com)",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    sub = p.add_subparsers(dest="cmd", required=True)

    pl = sub.add_parser("list-v2", help="list ALL v2 events and resolve every slug")
    pl.add_argument("--details", action="store_true",
                    help="also fetch full detail + tickets for every event")
    pl.set_defaults(func=cmd_list_v2)

    pv = sub.add_parser("v2", help="scrape specific v2 slug events, e.g. indo-comic")
    pv.add_argument("slugs", nargs="+", help="one or more event slugs")
    pv.set_defaults(func=cmd_v2)

    pe = sub.add_parser("extract-codes",
                        help="find invitation/compliment codes across all v2 events")
    pe.set_defaults(func=cmd_extract_codes)

    pw = sub.add_parser("watch",
                        help="detect NEW events / codes since last run (Telegram / desktop notif)")
    pw.add_argument("--notify", action="store_true",
                    help="also send a desktop notification (macOS/Linux)")
    pw.add_argument("--include-test", action="store_true",
                    help="also report internal/test events (default: ignore them)")
    pw.add_argument("--telegram-token", default="", dest="telegram_token",
                    help="Telegram bot token (or set env TELEGRAM_BOT_TOKEN)")
    pw.add_argument("--telegram-chat", default="", dest="telegram_chat",
                    help="Telegram chat id (or set env TELEGRAM_CHAT_ID)")
    pw.add_argument("--no-telegram", action="store_true",
                    help="do not send Telegram even if token/chat are configured")
    pw.add_argument("--test-telegram", action="store_true",
                    help="send a test Telegram message and exit")
    pw.set_defaults(func=cmd_watch)

    px = sub.add_parser("export-excel",
                        help="export all v2 events/tickets/codes to one .xlsx file")
    px.set_defaults(func=cmd_export_excel)

    ps = sub.add_parser("search", help="search legacy events by keyword")
    ps.add_argument("keyword")
    ps.add_argument("--pages", type=int, default=1)
    ps.add_argument("--per-page", type=int, default=40, dest="per_page")
    ps.set_defaults(func=cmd_search)

    pi = sub.add_parser("ids", help="fetch specific legacy numeric event ids")
    pi.add_argument("ids", nargs="+", type=int)
    pi.set_defaults(func=cmd_ids)

    pc = sub.add_parser("crawl", help="sweep keywords to harvest legacy event ids")
    pc.add_argument("--max-pages", type=int, default=25, dest="max_pages")
    pc.add_argument("--per-page", type=int, default=40, dest="per_page")
    pc.add_argument("--details", action="store_true",
                    help="also fetch full detail for each event")
    pc.set_defaults(func=cmd_crawl)

    args = p.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
